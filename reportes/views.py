from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.db.models import F, Sum
from django.http import HttpResponse
from django.shortcuts import render

from openpyxl import Workbook

from usuarios.decorators import solo_administrador
from productos.models import Producto
from compras.models import Compra, MovimientoInventario
from ventas.models import Venta, DetalleVenta


def obtener_valor(valor):
    return valor() if callable(valor) else valor


def formato_cop(valor):
    if valor is None:
        valor = 0

    try:
        valor = Decimal(valor)
    except Exception:
        valor = Decimal(0)

    return f"${valor:,.0f}".replace(",", ".")


@login_required
@solo_administrador
def dashboard_reportes(request):
    productos = Producto.objects.all()
    compras = Compra.objects.all().order_by("-fecha")
    ventas = Venta.objects.all().order_by("-fecha")

    total_productos = productos.count()
    total_compras_registradas = compras.count()
    total_ventas_registradas = ventas.count()

    productos_bajo_stock = Producto.objects.filter(
        stock__lte=F("stock_minimo")
    ).order_by("stock")

    total_bajo_stock = productos_bajo_stock.count()

    total_compras_dinero = Decimal(0)
    for compra in compras:
        total_compras_dinero += Decimal(obtener_valor(compra.total) or 0)

    total_ventas_dinero = Decimal(0)
    for venta in ventas:
        total_ventas_dinero += Decimal(obtener_valor(venta.total) or 0)

    productos_mas_vendidos = (
        DetalleVenta.objects
        .values("producto__nombre")
        .annotate(cantidad_vendida=Sum("cantidad"))
        .order_by("-cantidad_vendida")[:5]
    )

    ultimas_compras = compras[:5]
    ultimas_ventas = ventas[:5]

    ultimos_movimientos = (
        MovimientoInventario.objects
        .select_related("producto")
        .all()
        .order_by("-fecha")[:10]
    )

    contexto = {
        "total_productos": total_productos,
        "total_compras_registradas": total_compras_registradas,
        "total_ventas_registradas": total_ventas_registradas,
        "total_bajo_stock": total_bajo_stock,

        "total_compras_dinero_formateado": formato_cop(total_compras_dinero),
        "total_ventas_dinero_formateado": formato_cop(total_ventas_dinero),

        "productos_bajo_stock": productos_bajo_stock,
        "productos_mas_vendidos": productos_mas_vendidos,
        "ultimas_compras": ultimas_compras,
        "ultimas_ventas": ultimas_ventas,
        "ultimos_movimientos": ultimos_movimientos,
    }

    return render(request, "reportes/dashboard_reportes.html", contexto)


@login_required
@solo_administrador
def exportar_reporte_excel(request):
    wb = Workbook()

    ws = wb.active
    ws.title = "Resumen"

    productos = Producto.objects.all()
    compras = Compra.objects.all().order_by("-fecha")
    ventas = Venta.objects.all().order_by("-fecha")

    productos_bajo_stock = Producto.objects.filter(
        stock__lte=F("stock_minimo")
    )

    total_compras_dinero = Decimal(0)
    for compra in compras:
        total_compras_dinero += Decimal(obtener_valor(compra.total) or 0)

    total_ventas_dinero = Decimal(0)
    for venta in ventas:
        total_ventas_dinero += Decimal(obtener_valor(venta.total) or 0)

    ws.append(["REPORTE GENERAL DEL INVENTARIO"])
    ws.append([])
    ws.append(["Indicador", "Valor"])
    ws.append(["Total productos", productos.count()])
    ws.append(["Total compras registradas", compras.count()])
    ws.append(["Total ventas registradas", ventas.count()])
    ws.append(["Productos con bajo inventario", productos_bajo_stock.count()])
    ws.append(["Total compras COP", float(total_compras_dinero)])
    ws.append(["Total ventas COP", float(total_ventas_dinero)])

    ws_productos = wb.create_sheet("Productos")
    ws_productos.append([
        "Producto",
        "Categoría",
        "Proveedor",
        "Stock",
        "Stock mínimo",
        "Precio compra",
        "Precio venta",
    ])

    for producto in productos:
        ws_productos.append([
            producto.nombre,
            producto.categoria.nombre if producto.categoria else "",
            producto.proveedor.nombre if producto.proveedor else "",
            producto.stock,
            producto.stock_minimo,
            float(producto.precio_compra),
            float(producto.precio_venta),
        ])

    ws_compras = wb.create_sheet("Compras")
    ws_compras.append([
        "ID",
        "Proveedor",
        "Usuario",
        "Fecha",
        "Total",
        "Observación",
    ])

    for compra in compras:
        ws_compras.append([
            compra.id,
            compra.proveedor.nombre if compra.proveedor else "",
            compra.usuario.username if compra.usuario else "",
            compra.fecha.strftime("%d/%m/%Y %H:%M"),
            float(Decimal(obtener_valor(compra.total) or 0)),
            compra.observacion,
        ])

    ws_ventas = wb.create_sheet("Ventas")
    ws_ventas.append([
        "ID",
        "Cliente",
        "Usuario",
        "Fecha",
        "Total",
        "Observación",
    ])

    for venta in ventas:
        ws_ventas.append([
            venta.id,
            venta.cliente,
            venta.usuario.username if venta.usuario else "",
            venta.fecha.strftime("%d/%m/%Y %H:%M"),
            float(Decimal(obtener_valor(venta.total) or 0)),
            venta.observacion,
        ])

    ws_movimientos = wb.create_sheet("Movimientos")
    ws_movimientos.append([
        "ID",
        "Producto",
        "Tipo",
        "Cantidad",
        "Fecha",
        "Descripción",
    ])

    movimientos = MovimientoInventario.objects.select_related("producto").all().order_by("-fecha")

    for movimiento in movimientos:
        ws_movimientos.append([
            movimiento.id,
            movimiento.producto.nombre if movimiento.producto else "",
            movimiento.tipo,
            movimiento.cantidad,
            movimiento.fecha.strftime("%d/%m/%Y %H:%M"),
            movimiento.descripcion,
        ])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    response["Content-Disposition"] = 'attachment; filename="reporte_inventario.xlsx"'

    wb.save(response)

    return response 
     